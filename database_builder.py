# ============================================================
# database_builder.py - แปลงไฟล์ HTML ตารางคำทับศัพท์เป็นฐานข้อมูล
# ============================================================

import os
import sys
import re
import html
import sqlite3
import json
import csv
from datetime import datetime

# บังคับใช้ UTF-8 สำหรับ stdout บน Windows
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

SOURCE_HTML_PATH = r"C:\Users\NB_DELL\.gemini\antigravity\brain\4b8de4dd-ec8a-4d81-9eec-2d0116824a0c\.user_uploaded\media_1789014062063.html"
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
SQLITE_DB_PATH = os.path.join(DATA_DIR, "vocab.db")
CSV_PATH = os.path.join(DATA_DIR, "transliterations.csv")
ENGINE_CSV_PATH = os.path.join(DATA_DIR, "vocab_db.csv")
EXCEL_PATH = os.path.join(DATA_DIR, "transliterations.xlsx")
JSON_PATH = os.path.join(DATA_DIR, "transliterations.json")


def parse_html_table(html_path: str) -> list:
    """อ่านและสกัดแถวข้อมูลจากไฟล์ HTML Google Sheets"""
    if not os.path.exists(html_path):
        raise FileNotFoundError(f"ไม่พบไฟล์: {html_path}")

    with open(html_path, "r", encoding="utf-8") as f:
        content = f.read()

    tr_blocks = re.findall(r"<tr[^>]*>(.*?)</tr>", content, re.DOTALL)
    entries = []

    for tr in tr_blocks:
        tds = re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", tr, re.DOTALL)
        clean_tds = [html.unescape(re.sub(r"<[^>]+>", "", td)).strip() for td in tds]

        # โครงสร้างคอลัมน์:
        # clean_tds[0] = row header (เช่น 2, 3...)
        # clean_tds[1] = ลำดับที่ (1, 2, ...)
        # clean_tds[2] = ภาษาอังกฤษ
        # clean_tds[3] = คำทับศัพท์ภาษาไทย
        # clean_tds[4] = หมายเหตุ (ถ้ามี)
        if len(clean_tds) >= 4 and clean_tds[1].isdigit():
            num = int(clean_tds[1])
            en = clean_tds[2].strip()
            th = clean_tds[3].strip()
            note = clean_tds[4].strip() if len(clean_tds) > 4 else ""

            # ปรับแต่งคำผิดเล็กน้อยในหมายเหตุของราชบัณฑิตยสภา
            if "พจนานุกรม" in note:
                note = "พจนานุกรม"
            elif "ศัพท์บัญญัติ" in note or "ศััพท์บัญญัติ" in note:
                note = "ศัพท์บัญญัติ"
            elif "หลักนิยม" in note or "หลัักนิยม" in note or "หลักนิิยม" in note:
                note = "คำตามหลักนิยม"

            if en or th:
                entries.append({
                    "id": num,
                    "english_word": en,
                    "thai_word": th,
                    "note": note,
                    "category": "transliteration",
                })

    return entries


def build_sqlite_db(entries: list, db_path: str):
    """สร้างฐานข้อมูล SQLite พร้อม Table และ Index"""
    os.makedirs(os.path.dirname(db_path), exist_ok=True)
    conn = sqlite3.connect(db_path)
    cur = conn.cursor()

    cur.execute("""
        CREATE TABLE IF NOT EXISTS transliterations (
            id INTEGER PRIMARY KEY,
            english_word TEXT NOT NULL,
            thai_word TEXT NOT NULL,
            note TEXT DEFAULT '',
            category TEXT DEFAULT 'transliteration',
            is_active INTEGER DEFAULT 1,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    """)

    # ล้างข้อมูลเดิมถ้ามี เพื่อนำเข้าชุดใหม่ที่สมบูรณ์
    cur.execute("DELETE FROM transliterations")

    cur.executemany("""
        INSERT INTO transliterations (id, english_word, thai_word, note, category)
        VALUES (:id, :english_word, :thai_word, :note, :category)
    """, entries)

    # สร้าง Index เพื่อการสืบค้นรวดเร็วระดับ milliseconds
    cur.execute("CREATE INDEX IF NOT EXISTS idx_translit_en ON transliterations(english_word COLLATE NOCASE)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_translit_th ON transliterations(thai_word)")

    conn.commit()
    conn.close()
    print(f"✓ สร้าง SQLite DB สำเร็จ: {db_path} ({len(entries)} รายการ)")


def build_csv_files(entries: list, csv_path: str, engine_csv_path: str):
    """สร้างไฟล์ CSV สำหรับเปิดใน Excel (UTF-8 with BOM) และสำหรับ Engine"""
    # 1. CSV มาตรฐาน
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["ลำดับที่", "ภาษาอังกฤษ", "คำทับศัพท์ภาษาไทย", "หมายเหตุ"])
        for e in entries:
            writer.writerow([e["id"], e["english_word"], e["thai_word"], e["note"]])
    print(f"✓ สร้าง CSV มาตรฐานสำเร็จ: {csv_path}")

    # 2. CSV สำหรับ Engine / Google Sheets
    with open(engine_csv_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["incorrect_word", "correct_word", "note", "type"])
        for e in entries:
            # ใช้คำภาษาอังกฤษเป็น incorrect_word เพื่อแนะนำให้เขียนเป็นทับศัพท์ไทย
            writer.writerow([
                e["english_word"],
                e["thai_word"],
                f"คำทับศัพท์ทางการวุฒิสภา{' (' + e['note'] + ')' if e['note'] else ''}",
                "transliteration"
            ])
    print(f"✓ สร้าง Engine CSV สำเร็จ: {engine_csv_path}")


def build_excel_file(entries: list, xlsx_path: str):
    """สร้างไฟล์ Excel .xlsx สวยงามพร้อมจัดรูปแบบหัวตาราง"""
    wb = Workbook()
    ws = wb.active
    ws.title = "คำทับศัพท์วุฒิสภา"

    headers = ["ลำดับที่", "ภาษาอังกฤษ", "คำทับศัพท์ภาษาไทย", "หมายเหตุ"]
    ws.append(headers)

    # จัดรูปแบบหัวตาราง
    header_fill = PatternFill(start_color="1A237E", end_color="1A237E", fill_type="solid")
    header_font = Font(name="Sarabun", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0"),
    )

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    # ใส่ข้อมูล
    data_font = Font(name="Sarabun", size=10)
    for e in entries:
        row = [e["id"], e["english_word"], e["thai_word"], e["note"]]
        ws.append(row)

    # จัดความกว้างคอลัมน์
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 35
    ws.column_dimensions["C"].width = 35
    ws.column_dimensions["D"].width = 25

    # จัด alignment
    for row in ws.iter_rows(min_row=2, max_row=len(entries) + 1):
        row[0].alignment = Alignment(horizontal="center")
        for cell in row:
            cell.font = data_font
            cell.border = thin_border

    wb.save(xlsx_path)
    print(f"✓ สร้าง Excel สำเร็จ: {xlsx_path}")


def build_json_file(entries: list, json_path: str):
    """สร้างไฟล์ JSON"""
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(entries, f, ensure_ascii=False, indent=2)
    print(f"✓ สร้าง JSON สำเร็จ: {json_path}")


def main():
    print("🚀 กำลังแปลงไฟล์ HTML เป็นฐานข้อมูล...")
    os.makedirs(DATA_DIR, exist_ok=True)
    entries = parse_html_table(SOURCE_HTML_PATH)
    print(f"✓ สกัดข้อมูลได้ทั้งหมด {len(entries)} คำ")

    build_sqlite_db(entries, SQLITE_DB_PATH)
    build_csv_files(entries, CSV_PATH, ENGINE_CSV_PATH)
    build_excel_file(entries, EXCEL_PATH)
    build_json_file(entries, JSON_PATH)
    print("🎉 สร้างฐานข้อมูลครบทุกรูปแบบเรียบร้อยแล้ว!")


if __name__ == "__main__":
    main()
